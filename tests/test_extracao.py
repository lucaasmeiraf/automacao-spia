"""
Testes do ETL de anexos (`app/processing/extracao.py`) — arquivos montados no próprio teste.

A planilha imita a do Resumo do Projeto real (contrato 00 00493/2013, "Projeto Pavimento 493.xlsx").
"""
from __future__ import annotations

import io
import zipfile
from datetime import datetime

import openpyxl
import pytest

from app.processing.extracao import FormatoNaoSuportado, detectar_formato, extrair_texto


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo_Proj_PAVNOVO"
    ws.append(["ELABORAÇÃO DO PROJETO"])
    ws.append(["Empresa Responsável", None, "Construtora Continental - Sogel"])
    ws.append(["Ano de Elaboração", datetime(2015, 8, 1)])
    ws.append([None, None])  # linha vazia: some
    ws.append(["Extensão", 5.1, "Nº de Faixas", 2.0])
    ws.append(["Taxa de Cresc.", 0.035])
    ws["B6"].number_format = "0.0%"
    ws.append(["Texto com\nquebra de linha"])
    ws2 = wb.create_sheet("Pavimento")
    ws2.append(["Revestimento", "CBUQ", 5])
    wb.create_sheet("Vazia")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(*paginas: str) -> bytes:
    """PDF mínimo com texto (Helvetica), xref com offsets corretos."""
    objetos: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(paginas)))
    objetos.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(paginas)} >>".encode())
    fonte = 3 + 2 * len(paginas)
    for i, texto in enumerate(paginas):
        stream = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode("latin-1")
        objetos.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {fonte} 0 R >> >> >>".encode()
        )
        objetos.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objetos.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    saida = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, obj in enumerate(objetos, start=1):
        offsets.append(len(saida))
        saida += f"{n} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(saida)
    saida += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode()
    saida += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    saida += f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(saida)


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------

def test_xlsx_vira_texto_por_aba_e_linha():
    r = extrair_texto(_xlsx(), "Projeto.xlsx")

    assert r.formato == "xlsx"
    assert r.truncado is False and r.aviso == ""
    assert r.texto.startswith("### Aba: Resumo_Proj_PAVNOVO\nELABORAÇÃO DO PROJETO\n")
    assert "Empresa Responsável | Construtora Continental - Sogel" in r.texto  # célula vazia omitida
    assert "Ano de Elaboração | 2015-08-01" in r.texto
    assert "Extensão | 5.1 | Nº de Faixas | 2" in r.texto  # 2.0 → 2
    assert "Taxa de Cresc. | 3.5%" in r.texto  # formato % da célula
    assert "Texto com quebra de linha" in r.texto
    assert "### Aba: Pavimento\nRevestimento | CBUQ | 5" in r.texto
    assert "Vazia" not in r.texto  # aba sem conteúdo não aparece


def test_xlsx_respeita_o_limite():
    r = extrair_texto(_xlsx(), limite=40)
    assert r.truncado is True
    assert len(r.texto) <= 40


def test_xlsx_corrompido_e_formato_nao_suportado():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", "isto não é xml")
    with pytest.raises(FormatoNaoSuportado, match="planilha ilegível"):
        extrair_texto(buf.getvalue())


# ---------------------------------------------------------------------------
# pdf
# ---------------------------------------------------------------------------

def test_pdf_texto_por_pagina():
    r = extrair_texto(_pdf("Estrutura do pavimento CBUQ 5 cm", "Usina de asfalto"), "a.pdf")

    assert r.formato == "pdf"
    assert "### Página 1\nEstrutura do pavimento CBUQ 5 cm" in r.texto
    assert "### Página 2\nUsina de asfalto" in r.texto
    assert r.aviso == ""


def test_pdf_sem_texto_avisa_que_nao_ha_ocr():
    r = extrair_texto(_pdf(""))
    assert r.texto == ""
    assert "OCR" in r.aviso


def test_pdf_ilegivel():
    with pytest.raises(FormatoNaoSuportado):
        extrair_texto(b"%PDF-1.4\nlixo sem estrutura")


def test_pdf_respeita_o_limite():
    r = extrair_texto(_pdf("A" * 50, "B" * 50, "C" * 50), limite=60)
    assert r.truncado is True
    assert len(r.texto) <= 60
    assert "C" not in r.texto  # parou de ler antes da última página


# ---------------------------------------------------------------------------
# Detecção de formato
# ---------------------------------------------------------------------------

def test_formato_pelos_bytes_e_nao_pelo_nome():
    assert detectar_formato(_xlsx(), "enganoso.pdf") == "xlsx"
    assert detectar_formato(_pdf("x"), "enganoso.xlsx") == "pdf"


def test_csv_e_texto_puro():
    r = extrair_texto("km;valor\n324;5,1\n".encode("cp1252"), "dados.csv", "text/csv")
    assert r.formato == "texto"
    assert r.texto == "km;valor\n324;5,1"


@pytest.mark.parametrize("conteudo, trecho", [
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16, "formato antigo"),
    (b"\xff\xd8\xff\xe0imagem", "não suportado"),
])
def test_formatos_sem_suporte(conteudo, trecho):
    with pytest.raises(FormatoNaoSuportado, match=trecho):
        extrair_texto(conteudo, "arquivo.bin", "application/octet-stream")


def test_docx_ainda_nao_suportado():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", "<w:document/>")
    with pytest.raises(FormatoNaoSuportado, match="Word"):
        detectar_formato(buf.getvalue())
